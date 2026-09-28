"""CSV/XLSX table adapters that preserve the user's original structure."""

from __future__ import annotations

import codecs
import csv
import os
import tempfile
from abc import ABC, abstractmethod
from contextlib import contextmanager, suppress
from pathlib import Path
from typing import Any, Dict, Iterable, Iterator, List, Mapping, Optional, Sequence, Union


SheetName = Union[str, int]
DEFAULT_CSV_ENCODING = "utf-8-sig"
DEFAULT_CSV_DELIMITER = ","
DEFAULT_FORMULA_POLICY = "escape"


def _text_header(value: Any) -> str:
    return "" if value is None else str(value)


def _has_value(value: Any) -> bool:
    return value is not None and str(value).strip() != ""


def _validate_csv_options(encoding: str, delimiter: str) -> None:
    if not isinstance(encoding, str) or not encoding.strip():
        raise ValueError("csv_encoding must be a non-empty codec name.")
    try:
        codecs.lookup(encoding)
    except LookupError:
        raise ValueError(f"Unknown CSV encoding: {encoding!r}") from None
    if not isinstance(delimiter, str) or len(delimiter) != 1 or delimiter in "\r\n":
        raise ValueError("csv_delimiter must be exactly one non-newline character.")


def _validate_formula_policy(policy: str) -> None:
    if policy not in {"escape", "preserve"}:
        raise ValueError("formula_policy must be either 'escape' or 'preserve'.")


def _looks_like_formula(value: Any) -> bool:
    return isinstance(value, str) and value.lstrip(" \t\r\n").startswith(
        ("=", "+", "-", "@")
    )


def _csv_output_value(value: Any, formula_policy: str) -> Any:
    if formula_policy == "escape" and _looks_like_formula(value):
        return "'" + value
    return value


def _validate_file_headers(
    values: Sequence[Any],
    *,
    source: str,
    allow_trailing_empty: bool = False,
) -> List[str]:
    raw = list(values)
    if allow_trailing_empty:
        while raw and not _has_value(raw[-1]):
            raw.pop()
    if not raw:
        raise ValueError(f"{source} has no header row.")

    headers: List[str] = []
    seen = set()
    for column_index, value in enumerate(raw, start=1):
        if not _has_value(value):
            raise ValueError(
                f"{source} header at column {column_index} is empty."
            )
        name = _text_header(value)
        if name in seen:
            raise ValueError(f"Duplicate {source} column name: {name!r}")
        seen.add(name)
        headers.append(name)
    return headers


def _validate_output_headers(values: Sequence[str]) -> List[str]:
    if isinstance(values, (str, bytes)):
        raise TypeError("columns must be a sequence of strings, not a string.")

    headers: List[str] = []
    seen = set()
    for column_index, value in enumerate(values, start=1):
        if not isinstance(value, str):
            raise TypeError(
                "Column names must be strings; "
                f"column {column_index} has type {type(value).__name__}."
            )
        if not value.strip():
            raise ValueError(f"Output column {column_index} is empty.")
        if value in seen:
            raise ValueError(f"Duplicate output column name: {value!r}")
        seen.add(value)
        headers.append(value)
    if not headers:
        raise ValueError("At least one output column is required.")
    return headers


def _copy_output_row(
    record: Mapping[str, Any],
    *,
    row_number: int,
) -> Dict[str, Any]:
    row = dict(record)
    for name in row:
        if not isinstance(name, str):
            raise TypeError(
                "Row keys must be strings; "
                f"row {row_number} contains a {type(name).__name__} key."
            )
    return row


def _validate_output_row(
    record: Mapping[str, Any],
    *,
    headers: Sequence[str],
    row_number: int,
) -> Dict[str, Any]:
    row = _copy_output_row(record, row_number=row_number)
    allowed = set(headers)
    extra = [name for name in row if name not in allowed]
    if extra:
        raise ValueError(
            f"Row {row_number} contains columns not present in the output schema: "
            + ", ".join(repr(name) for name in extra)
        )
    return row


@contextmanager
def _atomic_output(path: Path) -> Iterator[Path]:
    """Yield a sibling temporary path and replace ``path`` only on success."""

    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(
        dir=str(path.parent),
        prefix=f".{path.stem}.",
        suffix=path.suffix,
    )
    os.close(descriptor)
    temporary = Path(temporary_name)
    try:
        yield temporary
        os.replace(temporary, path)
    finally:
        with suppress(OSError):
            temporary.unlink()


class Table(ABC):
    path: Path

    @property
    @abstractmethod
    def headers(self) -> Sequence[str]:
        raise NotImplementedError

    @property
    @abstractmethod
    def row_count(self) -> int:
        raise NotImplementedError

    @abstractmethod
    def ensure_column(self, name: str) -> None:
        raise NotImplementedError

    @abstractmethod
    def get(self, row_index: int, column: str) -> Any:
        raise NotImplementedError

    @abstractmethod
    def set(self, row_index: int, column: str, value: Any) -> None:
        raise NotImplementedError

    @abstractmethod
    def save(self, path: Path) -> None:
        raise NotImplementedError

    def close(self) -> None:
        """Release resources retained by this table."""

    def __enter__(self) -> "Table":
        return self

    def __exit__(self, exc_type: Any, exc_value: Any, traceback: Any) -> None:
        self.close()


class MemoryTable(Table):
    """A table backed by Python dictionaries instead of a file."""

    def __init__(self, rows: Iterable[Mapping[str, Any]]) -> None:
        self.path = Path("<memory>")
        self._rows = [
            _copy_output_row(row, row_number=row_number)
            for row_number, row in enumerate(rows, start=1)
        ]
        self._headers: List[str] = []
        for row in self._rows:
            for name in row:
                if not name.strip():
                    raise ValueError("Row keys cannot be empty.")
                if name not in self._headers:
                    self._headers.append(name)

    @property
    def headers(self) -> Sequence[str]:
        return tuple(self._headers)

    @property
    def row_count(self) -> int:
        return len(self._rows)

    def ensure_column(self, name: str) -> None:
        if name not in self._headers:
            self._headers.append(name)
        for row in self._rows:
            row.setdefault(name, "")

    def get(self, row_index: int, column: str) -> Any:
        return self._rows[row_index].get(column)

    def set(self, row_index: int, column: str, value: Any) -> None:
        self._rows[row_index][column] = value

    def save(self, path: Path) -> None:
        write_records(self._rows, path, fieldnames=self._headers)

    def records(self) -> List[Dict[str, Any]]:
        return [dict(row) for row in self._rows]


class CsvTable(Table):
    def __init__(
        self,
        path: Path,
        *,
        encoding: str = DEFAULT_CSV_ENCODING,
        delimiter: str = DEFAULT_CSV_DELIMITER,
        formula_policy: str = DEFAULT_FORMULA_POLICY,
    ) -> None:
        _validate_csv_options(encoding, delimiter)
        _validate_formula_policy(formula_policy)
        self.path = path
        self.encoding = encoding
        self.delimiter = delimiter
        self.formula_policy = formula_policy
        with path.open("r", encoding=encoding, newline="") as stream:
            reader = csv.DictReader(stream, delimiter=delimiter)
            if reader.fieldnames is None:
                raise ValueError(f"CSV file has no header row: {path}")
            self._headers = _validate_file_headers(reader.fieldnames, source="CSV")
            self._rows = []
            for row in reader:
                if None in row:
                    raise ValueError(
                        f"CSV row {reader.line_num} has more values than the header row."
                    )
                self._rows.append(
                    {header: row.get(header) for header in self._headers}
                )

    @property
    def headers(self) -> Sequence[str]:
        return tuple(self._headers)

    @property
    def row_count(self) -> int:
        return len(self._rows)

    def ensure_column(self, name: str) -> None:
        if name not in self._headers:
            self._headers.append(name)
            for row in self._rows:
                row[name] = ""

    def get(self, row_index: int, column: str) -> Any:
        return self._rows[row_index].get(column)

    def set(self, row_index: int, column: str, value: Any) -> None:
        self._rows[row_index][column] = value

    def save(self, path: Path) -> None:
        write_records(
            self._rows,
            path,
            fieldnames=self._headers,
            csv_encoding=self.encoding,
            csv_delimiter=self.delimiter,
            formula_policy=self.formula_policy,
        )


class XlsxTable(Table):
    def __init__(
        self,
        path: Path,
        sheet_name: SheetName,
        *,
        formula_policy: str = DEFAULT_FORMULA_POLICY,
    ) -> None:
        _validate_formula_policy(formula_policy)
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ImportError(
                "XLSX support requires openpyxl. "
                "Install aicorpusx with its dependencies."
            ) from exc

        self.path = path
        self.formula_policy = formula_policy
        self.workbook = load_workbook(path)
        try:
            if isinstance(sheet_name, int):
                try:
                    self.sheet = self.workbook.worksheets[sheet_name]
                except IndexError:
                    raise ValueError(
                        f"XLSX sheet index is out of range: {sheet_name}"
                    ) from None
            else:
                try:
                    self.sheet = self.workbook[str(sheet_name)]
                except KeyError:
                    raise ValueError(
                        f"XLSX sheet does not exist: {sheet_name!r}"
                    ) from None
            self._refresh_headers()
        except BaseException:
            self.workbook.close()
            raise

    def _refresh_headers(self) -> None:
        if self.formula_policy == "escape":
            for cell in self.sheet[1]:
                if _looks_like_formula(cell.value):
                    cell.data_type = "s"
        headers = _validate_file_headers(
            [cell.value for cell in self.sheet[1]],
            source="XLSX",
            allow_trailing_empty=True,
        )
        self._columns = {
            name: column_index for column_index, name in enumerate(headers, start=1)
        }
        if self.sheet.max_column > len(headers):
            extra_rows = self.sheet.iter_rows(
                min_row=2,
                min_col=len(headers) + 1,
                max_col=self.sheet.max_column,
                values_only=True,
            )
            for row_number, values in enumerate(extra_rows, start=2):
                if any(_has_value(value) for value in values):
                    raise ValueError(
                        f"XLSX row {row_number} has values beyond the header columns."
                    )

    @property
    def headers(self) -> Sequence[str]:
        return tuple(self._columns)

    @property
    def row_count(self) -> int:
        return max(0, self.sheet.max_row - 1)

    def ensure_column(self, name: str) -> None:
        if name not in self._columns:
            column = max(self._columns.values(), default=0) + 1
            cell = self.sheet.cell(row=1, column=column, value=name)
            if self.formula_policy == "escape" and _looks_like_formula(name):
                cell.data_type = "s"
            self._columns[name] = column

    def get(self, row_index: int, column: str) -> Any:
        return self.sheet.cell(row=row_index + 2, column=self._columns[column]).value

    def set(self, row_index: int, column: str, value: Any) -> None:
        cell = self.sheet.cell(
            row=row_index + 2,
            column=self._columns[column],
            value=value,
        )
        if self.formula_policy == "escape" and _looks_like_formula(value):
            cell.data_type = "s"

    def save(self, path: Path) -> None:
        with _atomic_output(path) as temporary:
            self.workbook.save(temporary)

    def close(self) -> None:
        self.workbook.close()


def open_table(
    path: Path,
    sheet_name: SheetName = 0,
    *,
    csv_encoding: str = DEFAULT_CSV_ENCODING,
    csv_delimiter: str = DEFAULT_CSV_DELIMITER,
    formula_policy: str = DEFAULT_FORMULA_POLICY,
) -> Table:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return CsvTable(
            path,
            encoding=csv_encoding,
            delimiter=csv_delimiter,
            formula_policy=formula_policy,
        )
    if suffix == ".xlsx":
        return XlsxTable(path, sheet_name, formula_policy=formula_policy)
    raise ValueError("Only .csv and .xlsx files are supported.")


def read_records(
    path: Path,
    sheet_name: SheetName = 0,
    *,
    csv_encoding: str = DEFAULT_CSV_ENCODING,
    csv_delimiter: str = DEFAULT_CSV_DELIMITER,
) -> tuple[List[str], List[Dict[str, Any]]]:
    table = open_table(
        path,
        sheet_name,
        csv_encoding=csv_encoding,
        csv_delimiter=csv_delimiter,
    )
    try:
        headers = list(table.headers)
        rows: List[Dict[str, Any]] = []
        for row_index in range(table.row_count):
            rows.append({header: table.get(row_index, header) for header in headers})
        return headers, rows
    finally:
        table.close()


def read_headers(
    path: Path,
    sheet_name: SheetName = 0,
    *,
    csv_encoding: str = DEFAULT_CSV_ENCODING,
    csv_delimiter: str = DEFAULT_CSV_DELIMITER,
    data_only: bool = False,
) -> List[str]:
    """Read and validate a table header without requiring a data row."""

    suffix = path.suffix.lower()
    if suffix == ".csv":
        _validate_csv_options(csv_encoding, csv_delimiter)
        with path.open("r", encoding=csv_encoding, newline="") as stream:
            reader = csv.reader(stream, delimiter=csv_delimiter)
            try:
                values = next(reader)
            except StopIteration:
                raise ValueError(f"CSV file has no header row: {path}") from None
        return _validate_file_headers(values, source="CSV")
    if suffix != ".xlsx":
        raise ValueError("Only .csv and .xlsx files are supported.")

    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ImportError(
            "XLSX support requires openpyxl. Install aicorpusx with its dependencies."
        ) from exc

    workbook = load_workbook(path, read_only=True, data_only=data_only)
    try:
        if isinstance(sheet_name, int):
            try:
                sheet = workbook.worksheets[sheet_name]
            except IndexError:
                raise ValueError(
                    f"XLSX sheet index is out of range: {sheet_name}"
                ) from None
        else:
            try:
                sheet = workbook[str(sheet_name)]
            except KeyError:
                raise ValueError(
                    f"XLSX sheet does not exist: {sheet_name!r}"
                ) from None
        rows = sheet.iter_rows(values_only=True)
        try:
            values = next(rows)
        except StopIteration:
            raise ValueError("XLSX has no header row.") from None
        return _validate_file_headers(
            values,
            source="XLSX",
            allow_trailing_empty=True,
        )
    finally:
        workbook.close()


def iter_records(
    path: Path,
    sheet_name: SheetName = 0,
    *,
    csv_encoding: str = DEFAULT_CSV_ENCODING,
    csv_delimiter: str = DEFAULT_CSV_DELIMITER,
    data_only: bool = False,
) -> Iterator[Dict[str, Any]]:
    """Yield CSV/XLSX records without loading the complete file into memory."""

    suffix = path.suffix.lower()
    if suffix == ".csv":
        _validate_csv_options(csv_encoding, csv_delimiter)
        with path.open("r", encoding=csv_encoding, newline="") as stream:
            reader = csv.DictReader(stream, delimiter=csv_delimiter)
            if reader.fieldnames is None:
                raise ValueError(f"CSV file has no header row: {path}")
            headers = _validate_file_headers(reader.fieldnames, source="CSV")
            for row in reader:
                if None in row:
                    raise ValueError(
                        f"CSV row {reader.line_num} has more values than the header row."
                    )
                yield {header: row.get(header) for header in headers}
        return
    if suffix != ".xlsx":
        raise ValueError("Only .csv and .xlsx files are supported.")

    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ImportError(
            "XLSX support requires openpyxl. "
            "Install aicorpusx with its dependencies."
        ) from exc

    workbook = load_workbook(path, read_only=True, data_only=data_only)
    try:
        if isinstance(sheet_name, int):
            try:
                sheet = workbook.worksheets[sheet_name]
            except IndexError:
                raise ValueError(f"XLSX sheet index is out of range: {sheet_name}") from None
        else:
            try:
                sheet = workbook[str(sheet_name)]
            except KeyError:
                raise ValueError(f"XLSX sheet does not exist: {sheet_name!r}") from None
        rows = sheet.iter_rows(values_only=True)
        try:
            first = next(rows)
        except StopIteration:
            raise ValueError("XLSX has no header row.") from None
        headers = _validate_file_headers(
            first,
            source="XLSX",
            allow_trailing_empty=True,
        )
        for row_number, values in enumerate(rows, start=2):
            if any(_has_value(value) for value in values[len(headers) :]):
                raise ValueError(
                    f"XLSX row {row_number} has values beyond the header columns."
                )
            yield {
                name: values[index] if index < len(values) else None
                for index, name in enumerate(headers)
            }
    finally:
        workbook.close()


def _write_records_from_iterator(
    iterator: Iterator[Mapping[str, Any]],
    path: Path,
    *,
    sheet_name: str = "Sheet1",
    fieldnames: Optional[Sequence[str]] = None,
    csv_encoding: str = DEFAULT_CSV_ENCODING,
    csv_delimiter: str = DEFAULT_CSV_DELIMITER,
    formula_policy: str = DEFAULT_FORMULA_POLICY,
) -> Path:
    suffix = path.suffix.lower()
    if suffix not in {".csv", ".xlsx"}:
        raise ValueError("Only .csv and .xlsx files are supported.")
    _validate_formula_policy(formula_policy)
    if suffix == ".csv":
        _validate_csv_options(csv_encoding, csv_delimiter)
    try:
        first_record = next(iterator)
    except StopIteration:
        first_record = None
    if fieldnames is None and first_record is None:
        raise ValueError("columns must be provided when writing an empty row iterable.")
    first_row = (
        _copy_output_row(first_record, row_number=1)
        if first_record is not None
        else None
    )
    header_values = fieldnames if fieldnames is not None else list(first_row or {})
    headers = _validate_output_headers(header_values)

    def checked_rows() -> Iterator[Dict[str, Any]]:
        row_number = 1
        if first_row is not None:
            yield _validate_output_row(
                first_row,
                headers=headers,
                row_number=row_number,
            )
        for row_number, record in enumerate(iterator, start=2):
            yield _validate_output_row(
                record,
                headers=headers,
                row_number=row_number,
            )

    with _atomic_output(path) as temporary:
        if suffix == ".csv":
            with temporary.open("w", encoding=csv_encoding, newline="") as stream:
                writer = csv.DictWriter(
                    stream,
                    fieldnames=headers,
                    delimiter=csv_delimiter,
                )
                writer.writerow(
                    {
                        name: _csv_output_value(name, formula_policy)
                        for name in headers
                    }
                )
                for row in checked_rows():
                    writer.writerow(
                        {
                            name: _csv_output_value(row.get(name), formula_policy)
                            for name in headers
                        }
                    )
        else:
            try:
                from openpyxl import Workbook
            except ImportError as exc:
                raise ImportError(
                    "XLSX support requires openpyxl. "
                    "Install aicorpusx with its dependencies."
                ) from exc
            workbook = Workbook(write_only=True)
            try:
                from openpyxl.cell import WriteOnlyCell

                sheet = workbook.create_sheet(title=str(sheet_name))
                header_values = []
                for name in headers:
                    if formula_policy == "escape" and _looks_like_formula(name):
                        cell = WriteOnlyCell(sheet, value=name)
                        cell.data_type = "s"
                        header_values.append(cell)
                    else:
                        header_values.append(name)
                sheet.append(header_values)
                try:
                    for row in checked_rows():
                        values = []
                        for name in headers:
                            value = row.get(name)
                            if formula_policy == "escape" and _looks_like_formula(value):
                                cell = WriteOnlyCell(sheet, value=value)
                                cell.data_type = "s"
                                values.append(cell)
                            else:
                                values.append(value)
                        sheet.append(values)
                except BaseException:
                    with suppress(Exception):
                        sheet.close()
                    raise
                workbook.save(temporary)
            finally:
                workbook.close()
    return path


def write_records(
    records: Iterable[Mapping[str, Any]],
    path: Path,
    *,
    sheet_name: str = "Sheet1",
    fieldnames: Optional[Sequence[str]] = None,
    csv_encoding: str = DEFAULT_CSV_ENCODING,
    csv_delimiter: str = DEFAULT_CSV_DELIMITER,
    formula_policy: str = DEFAULT_FORMULA_POLICY,
) -> Path:
    """Write records incrementally to CSV or XLSX and close input generators."""

    iterator = iter(records)
    try:
        return _write_records_from_iterator(
            iterator,
            path,
            sheet_name=sheet_name,
            fieldnames=fieldnames,
            csv_encoding=csv_encoding,
            csv_delimiter=csv_delimiter,
            formula_policy=formula_policy,
        )
    finally:
        close = getattr(iterator, "close", None)
        if callable(close):
            with suppress(Exception):
                close()


def is_filled(value: Any) -> bool:
    return _has_value(value)
